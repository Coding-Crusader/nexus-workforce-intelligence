"""
NEXUS Local Resume Parser
100% Local, privacy-preserving resume skill extractor for PDF and DOCX files.
Does NOT transmit any data externally.
"""

import io
import re
from typing import Dict, Any, List, Optional
import pypdf
import docx

from backend.data_loader import get_skill_aliases
from backend.skill_analysis import SkillAnalyzer

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extracts raw text from PDF bytes locally using pypdf."""
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        text_chunks = []
        for page in reader.pages:
            txt = page.extract_text()
            if txt:
                text_chunks.append(txt)
        return "\n".join(text_chunks)
    except Exception:
        try:
            return file_bytes.decode("utf-8", errors="ignore")
        except Exception:
            return ""

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extracts raw text from DOCX bytes locally using python-docx."""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        text_chunks = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        text_chunks.append(cell.text)
        return "\n".join(text_chunks)
    except Exception:
        try:
            return file_bytes.decode("utf-8", errors="ignore")
        except Exception:
            return ""

def parse_experience_from_text(text: str) -> float:
    """Heuristic extraction of total years of experience from resume text."""
    patterns = [
        r"(\d+(?:\.\d+)?)\+?\s*(?:years|yrs)\s+(?:of\s+)?experience",
        r"experience\s*:\s*(\d+(?:\.\d+)?)\+?\s*(?:years|yrs)",
        r"total\s+experience\s*:\s*(\d+(?:\.\d+)?)"
    ]
    for pat in patterns:
        m = re.search(pat, text, re.IGNORECASE)
        if m:
            try:
                val = float(m.group(1))
                if 0.0 <= val <= 40.0:
                    return val
            except ValueError:
                pass
    return 2.0  # default reasonable fallback

def extract_capabilities_from_resume(
    file_bytes: Any,
    filename: Any = "resume.pdf"
) -> Dict[str, Any]:
    """
    Locally parses resume bytes, identifies skills matching canonical taxonomy,
    and infers estimated proficiency scores.
    """
    # Defensive argument swap if caller passed (filename, file_bytes)
    if isinstance(file_bytes, str) and isinstance(filename, (bytes, bytearray)):
        file_bytes, filename = filename, file_bytes

    if not isinstance(filename, str):
        filename = str(filename)
    if not isinstance(file_bytes, (bytes, bytearray)):
        file_bytes = bytes(file_bytes) if file_bytes is not None else b""

    fname_lower = filename.lower()
    if fname_lower.endswith(".pdf"):
        text = extract_text_from_pdf(file_bytes)
    elif fname_lower.endswith(".docx") or fname_lower.endswith(".doc"):
        text = extract_text_from_docx(file_bytes)
    else:
        text = file_bytes.decode("utf-8", errors="ignore")

    text_lower = text.lower()
    aliases = get_skill_aliases()
    detected_skills = {}
    skill_counts = {}

    for canonical_key, variants in aliases.items():
        count = 0
        all_terms = [canonical_key.replace("_", " ")] + variants
        for term in all_terms:
            # Word boundary regex for term
            term_clean = re.escape(term.lower().strip())
            matches = len(re.findall(rf"\b{term_clean}\b", text_lower))
            count += matches

        skill_counts[canonical_key] = count
        # Strict Evidence Heuristic: Unverified resume text mentions
        # 0 count -> 0.20 unverified baseline
        # 1 count -> 0.35 foundational mention
        # 2 counts -> 0.48 developing knowledge
        # 3 counts -> 0.60 documented competency
        # 4-5 counts -> 0.72 recurring project implementation
        # 6+ counts -> 0.82 deep professional specialization
        if count == 0:
            score = 0.20
        elif count == 1:
            score = 0.35
        elif count == 2:
            score = 0.48
        elif count == 3:
            score = 0.60
        elif count <= 5:
            score = 0.72
        else:
            score = 0.82
        detected_skills[canonical_key] = score

    exp_years = parse_experience_from_text(text)

    return {
        "filename": filename,
        "estimated_experience_years": exp_years,
        "capabilities": detected_skills,
        "raw_mention_counts": skill_counts,
        "text_sample": text[:800] + ("..." if len(text) > 800 else "")
    }
